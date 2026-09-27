"""Compartimento B — integrita' dell'I/O: export, cache, temporanei, pubblicazione.

Copre B01, B10, B11, B12, R02, R03 e M05. B08/N06 stanno in
`test_decomposition_integrita.py`, accanto alle altre regressioni della cache.

Regola dei test di questo file: dove il contratto riguarda un file, il controllo
si fa sul FILE RILETTO da disco, non sull'oggetto in memoria.
"""
import os
import pathlib
import subprocess
import sys
import types
import urllib.parse
import urllib.request

import pytest

from gioco27.core import pdfmerge
from gioco27.core.algebra import EXCEL_MAX_CELL_CHARS, scrivi_excel
from gioco27.core.parallel import ExportTooLarge
from gioco27.gui import analysis_tab as at
from gioco27.gui import export_group_dialog as egd

RADICE = pathlib.Path(__file__).resolve().parents[1]
FOGLI_GREZZI = ["Dati grezzi", "Analisi della molteplicità", "Simbolica -> Perm"]


# ───────────────────────────── armamentario ─────────────────────────────────

class _Var:
    def __init__(self, v=""):
        self._v = v

    def get(self):
        return self._v

    def set(self, v):
        self._v = v


class _Analisi(at.AnalysisTabMixin):
    """Il mixin reale con lo stretto necessario intorno: nessuna finestra."""

    def __init__(self, righe=None, risultati=None):
        self._analisi_righe_raw = righe or [{"Stage0": "x"}]
        self._analisi_risultati = risultati or []
        self._analisi_status = _Var()
        self._closing = False
        self.progress = {}

    def _ui(self, fn):
        fn()

    def update_idletasks(self):
        pass


@pytest.fixture
def senza_dialoghi(monkeypatch):
    """Sostituisce filedialog/messagebox del tab Analisi; raccoglie gli errori."""
    errori = []
    monkeypatch.setattr(at, "messagebox", types.SimpleNamespace(
        showerror=lambda *a, **k: errori.append(a),
        showinfo=lambda *a, **k: None,
        showwarning=lambda *a, **k: None,
        askyesno=lambda *a, **k: True))
    return errori


def _apri(path):
    openpyxl = pytest.importorskip("openpyxl")
    return openpyxl.load_workbook(path)


def _esporta_excel_grezzo(risultati, path, monkeypatch):
    monkeypatch.setattr(at, "filedialog", types.SimpleNamespace(
        asksaveasfilename=lambda **k: str(path),
        askopenfilename=lambda **k: "", askdirectory=lambda **k: ""))
    tab = _Analisi(risultati=risultati)
    tab._export_analisi_raw_excel()
    return tab


def _risultato(sequenze, perm="[0,1,2]"):
    return {"n_sim": len(sequenze), "perm_str": perm, "simboliche": list(sequenze)}


# ─────────────────────────────── B01 ────────────────────────────────────────

def test_b01_centinaia_di_sequenze_sono_tutte_nel_file(tmp_path, monkeypatch,
                                                       senza_dialoghi):
    pytest.importorskip("openpyxl")
    sequenze = [f"(SCD_U x CSD_U x DCS_U) o MSC #{i:04d} " + "=" * 80
                for i in range(400)]
    path = tmp_path / "grezzi.xlsx"
    tab = _esporta_excel_grezzo([_risultato(sequenze)], path, monkeypatch)

    assert senza_dialoghi == [], "l'export non deve segnalare errori"
    assert tab._analisi_status.get().startswith("✓")
    wb = _apri(path)
    try:
        assert wb.sheetnames == FOGLI_GREZZI
        # il riepilogo supera il limite: diventa esplicito, non troncato di nascosto
        riepilogo = wb["Analisi della molteplicità"].cell(row=2, column=3).value
        assert len(riepilogo) <= EXCEL_MAX_CELL_CHARS
        assert str(len(sequenze)) in riepilogo
        assert "Simbolica -> Perm" in riepilogo
        # il dettaglio e' completo e ricostruibile
        dettaglio = [r[0] for r in wb["Simbolica -> Perm"].iter_rows(
            min_row=2, max_col=1, values_only=True)]
        assert dettaglio == sequenze
    finally:
        wb.close()


@pytest.mark.parametrize("lunghezza", [EXCEL_MAX_CELL_CHARS - 1,
                                       EXCEL_MAX_CELL_CHARS,
                                       EXCEL_MAX_CELL_CHARS + 1])
def test_b01_confini_del_limite_di_cella(lunghezza, tmp_path, monkeypatch,
                                         senza_dialoghi):
    """32.766 / 32.767 / 32.768 caratteri: sotto soglia il riepilogo resta intero."""
    pytest.importorskip("openpyxl")
    # due sequenze separate da "\n": la lunghezza totale e' quella richiesta
    prima = "A" * (lunghezza // 2)
    seconda = "B" * (lunghezza - len(prima) - 1)
    sequenze = [prima, seconda]
    assert len("\n".join(sequenze)) == lunghezza

    path = tmp_path / f"grezzi_{lunghezza}.xlsx"
    _esporta_excel_grezzo([_risultato(sequenze)], path, monkeypatch)
    assert senza_dialoghi == []
    wb = _apri(path)
    try:
        riepilogo = wb["Analisi della molteplicità"].cell(row=2, column=3).value
        if lunghezza <= EXCEL_MAX_CELL_CHARS:
            assert riepilogo == "\n".join(sequenze)
        else:
            assert riepilogo != "\n".join(sequenze)
            assert len(riepilogo) <= EXCEL_MAX_CELL_CHARS
        # in ogni caso il dettaglio contiene entrambe le sequenze, intere
        dettaglio = [r[0] for r in wb["Simbolica -> Perm"].iter_rows(
            min_row=2, max_col=1, values_only=True)]
        assert dettaglio == sequenze
    finally:
        wb.close()


def test_b01_schema_stabile_ed_etichette_umane_corrette(
    tmp_path, monkeypatch, senza_dialoghi,
):
    """Le colonne di schema restano stabili; le etichette umane sono italiane."""
    pytest.importorskip("openpyxl")
    path = tmp_path / "g.xlsx"
    righe = [{"Stage0": "s0", "Stage1": "s1", "Stage2": "s2", "A0": "a0",
              "A1": "a1", "A2": "a2", "T_simbolica": "sim",
              "T_permutazione": "[0]"}]
    monkeypatch.setattr(at, "filedialog", types.SimpleNamespace(
        asksaveasfilename=lambda **k: str(path)))
    tab = _Analisi(righe=righe, risultati=[_risultato(["x", "y"])])
    tab._export_analisi_raw_excel()
    wb = _apri(path)
    try:
        assert [c.value for c in wb["Dati grezzi"][1]] == [
            "#", "Stage0", "Stage1", "Stage2", "A0", "A1", "A2",
            "T_simbolica", "T_permutazione"]
        assert [c.value for c in wb["Analisi della molteplicità"][1]] == [
            "Molteplicità", "T_permutazione", "Sequenze di stadi distinte"]
        assert [c.value for c in wb["Simbolica -> Perm"][1]] == [
            "T_simbolica", "T_permutazione", "n_sim_distinte"]
    finally:
        wb.close()


# ─────────────────────────────── B10 ────────────────────────────────────────

def _fonts_puliti(monkeypatch):
    from gioco27.core import detail_pdf
    monkeypatch.setattr(detail_pdf, "_FONTS_READY", None)
    return detail_pdf


def test_b10_registrazione_fallita_da_fallback_reale(monkeypatch):
    pytest.importorskip("reportlab")
    from reportlab.pdfbase import pdfmetrics
    dp = _fonts_puliti(monkeypatch)

    def esplode(*a, **k):
        raise ValueError("registrazione impossibile")

    monkeypatch.setattr(pdfmetrics, "registerFont", esplode)
    assert dp._ensure_fonts() == ("Helvetica", "Helvetica-Bold")


def test_b10_registrazione_parziale_non_mischia_i_font(monkeypatch):
    """Il primo font passa, il secondo no: non si restituisce una coppia mista."""
    pytest.importorskip("reportlab")
    from reportlab.pdfbase import pdfmetrics
    dp = _fonts_puliti(monkeypatch)
    chiamate = {"n": 0}
    originale = pdfmetrics.registerFont

    def a_meta(font):
        chiamate["n"] += 1
        if chiamate["n"] == 1:
            return originale(font)
        raise ValueError("secondo font non registrabile")

    monkeypatch.setattr(pdfmetrics, "registerFont", a_meta)
    assert dp._ensure_fonts() == ("Helvetica", "Helvetica-Bold")


def test_b10_font_mancanti(monkeypatch, tmp_path):
    pytest.importorskip("reportlab")
    dp = _fonts_puliti(monkeypatch)
    monkeypatch.setattr(dp, "_ASSETS", str(tmp_path))
    assert dp._ensure_fonts() == ("Helvetica", "Helvetica-Bold")


def test_b10_font_corrotti(monkeypatch, tmp_path):
    pytest.importorskip("reportlab")
    dp = _fonts_puliti(monkeypatch)
    for nome in ("DejaVuSans.ttf", "DejaVuSans-Bold.ttf"):
        (tmp_path / nome).write_bytes(b"non sono un font")
    monkeypatch.setattr(dp, "_ASSETS", str(tmp_path))
    assert dp._ensure_fonts() == ("Helvetica", "Helvetica-Bold")


PROGRAMMA_FONT = """
import sys
sys.path.insert(0, {radice!r})
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfgen import canvas
from gioco27.core import detail_pdf

originale = pdfmetrics.registerFont
if {rompi!r}:
    def esplode(*a, **k):
        raise ValueError("registrazione impossibile")
    pdfmetrics.registerFont = esplode

regular, bold = detail_pdf._ensure_fonts()
# Il guasto riguardava la registrazione dei DejaVu: ripristinato l'originale,
# i font base restano quelli che ReportLab sa registrare da se'.
pdfmetrics.registerFont = originale
print("FONT", regular, bold)
# la prova vera: i nomi restituiti devono essere usabili davvero
c = canvas.Canvas({pdf!r})
for nome in (regular, bold):
    c.setFont(nome, 9)
    c.drawString(20, 20, "A\\u2665 27")
c.save()
print("PDF OK")
"""


@pytest.mark.parametrize("rompi, atteso", [(False, "DV DVB"),
                                           (True, "Helvetica Helvetica-Bold")])
def test_b10_processo_appena_avviato(rompi, atteso, tmp_path):
    """In un processo NUOVO: nessun font gia' registrato da un test precedente.

    Il PDF viene davvero disegnato con i nomi restituiti: un fallback che non
    si puo' usare farebbe fallire `setFont`.
    """
    pytest.importorskip("reportlab")
    pdf = tmp_path / "prova.pdf"
    res = subprocess.run(
        [sys.executable, "-c", PROGRAMMA_FONT.format(
            radice=str(RADICE), rompi=rompi, pdf=str(pdf))],
        capture_output=True, text=True, timeout=120,
        env=dict(os.environ, PYTHONDONTWRITEBYTECODE="1"))
    assert res.returncode == 0, res.stderr
    assert f"FONT {atteso}" in res.stdout
    assert "PDF OK" in res.stdout
    assert pdf.exists() and pdf.stat().st_size > 0


# ─────────────────────────────── B11 ────────────────────────────────────────

class _Sentinella(Exception):
    """Se viene sollevata, l'enumerazione e' partita: il preflight e' mancato."""


def test_b11_write_csv_rifiuta_prima_di_enumerare(tmp_path, monkeypatch):
    from gioco27.core import combinations, permutations

    def mai(_filters):
        raise _Sentinella("iter_combinations_ex non doveva essere chiamata")

    monkeypatch.setattr(combinations, "iter_combinations_ex", mai)
    liberi = [{"p0": "*", "p1": "*", "p2": "*",
               "j0": "*", "j1": "*", "j2": "*"} for _ in range(3)]
    assert combinations.count_combinations_ex(liberi) == 5_159_780_352
    dest = tmp_path / "enorme.csv"
    with pytest.raises(ExportTooLarge):
        permutations.write_csv(str(dest), liberi)
    assert not dest.exists(), "nessun file, nemmeno con la sola intestazione"
    assert list(tmp_path.iterdir()) == [], "nessun temporaneo lasciato in giro"


def test_b11_export_piccolo_resta_possibile(tmp_path):
    from gioco27.core.permutations import write_csv

    fissi = [{"p0": "SCD_U", "p1": "SCD_U", "p2": "SCD_U",
              "j0": "I_3", "j1": "I_3", "j2": "*"} for _ in range(3)]
    dest = tmp_path / "piccolo.csv"
    n = write_csv(str(dest), fissi)
    assert n == 8 and dest.exists()
    assert len(dest.read_text(encoding="utf-8").strip().splitlines()) == n + 1


def test_b11_tutte_le_rotte_pubbliche_hanno_il_preflight():
    """Nessun entry point di export massivo senza check_export_size."""
    import ast

    rotte = {
        # G2: le due rotte CSV sono passate a core/export_combinazioni.py
        # insieme al resto dell'orchestrazione degli export. Il preflight
        # e' lo stesso, nello stesso punto: prima di aprire il file.
        "gioco27/core/export_combinazioni.py": ["write_csv",
                                                "write_csv_parallel"],
        "gioco27/core/combinations.py": ["generate_pdf", "generate_pdf_ex",
                                         "_pdf_parallel"],
    }
    for rel, funzioni in rotte.items():
        albero = ast.parse((RADICE / rel).read_text(encoding="utf-8"))
        for nome in funzioni:
            fn = next(n for n in ast.walk(albero)
                      if isinstance(n, ast.FunctionDef) and n.name == nome)
            chiamate = {n.func.id for n in ast.walk(fn)
                        if isinstance(n, ast.Call) and isinstance(n.func, ast.Name)}
            assert "check_export_size" in chiamate, f"{rel}:{nome}"


# ─────────────────────────────── B12 ────────────────────────────────────────

class _FintoDialogo:
    """Il minimo per esercitare _content/_export_all del dialogo reale."""

    _PREVIEW_MAX = 50

    def __init__(self, items, cartella, scelti=None):
        self._items = items
        self._cache = {}
        self._cartella = cartella
        self._opts = {k: _Var(True if scelti is None else k in scelti)
                      for k, *_ in items}
        self._tabs = {k: _TestoFinto() for k, *_ in items}

    _content = egd._PreviewExportDialog._content
    _dimentica = egd._PreviewExportDialog._dimentica
    _refresh_preview = egd._PreviewExportDialog._refresh_preview
    _export_all = egd._PreviewExportDialog._export_all


class _TestoFinto:
    def __init__(self):
        self.text = ""

    def configure(self, **kw):
        pass

    def delete(self, *a):
        self.text = ""

    def insert(self, _dove, s):
        self.text += s


@pytest.fixture
def dialogo_senza_finestre(monkeypatch, tmp_path):
    mostrati = []
    monkeypatch.setattr(egd, "messagebox", types.SimpleNamespace(
        showinfo=lambda t, m, **k: mostrati.append(("info", t, m)),
        showerror=lambda t, m, **k: mostrati.append(("error", t, m))))
    monkeypatch.setattr(egd, "filedialog", types.SimpleNamespace(
        askdirectory=lambda **k: str(tmp_path)))
    return mostrati


def test_b12_generatore_fallito_non_produce_file(tmp_path, dialogo_senza_finestre):
    items = [("buono", "", "buono.svg", lambda: "<svg/>"),
             ("rotto", "", "rotto.svg", lambda: (_ for _ in ()).throw(
                 RuntimeError("generatore guasto")))]
    d = _FintoDialogo(items, tmp_path)
    d._export_all()

    assert (tmp_path / "buono.svg").read_text(encoding="utf-8") == "<svg/>"
    assert not (tmp_path / "rotto.svg").exists(), "nessun .svg falso"
    tipi = [t for t, *_ in dialogo_senza_finestre]
    assert "error" in tipi and "info" in tipi
    completato = [m for t, _ti, m in dialogo_senza_finestre if t == "info"][0]
    assert "buono.svg" in completato and "rotto.svg" not in completato


def test_b12_errore_non_e_contenuto():
    items = [("rotto", "", "rotto.tex", lambda: (_ for _ in ()).throw(
        ValueError("boom")))]
    d = _FintoDialogo(items, None)
    esito = d._content("rotto", items[0][3])
    assert esito.ok is False and esito.testo == "" and esito.errore == "boom"
    d._refresh_preview()
    assert "boom" in d._tabs["rotto"].text        # visibile solo in anteprima


def test_b12_ritentativo_dopo_correzione(tmp_path, dialogo_senza_finestre):
    stato = {"rotto": True}

    def generatore():
        if stato["rotto"]:
            raise RuntimeError("non ancora")
        return "contenuto valido"

    items = [("k", "", "k.tex", generatore)]
    d = _FintoDialogo(items, tmp_path)
    d._export_all()
    assert not (tmp_path / "k.tex").exists()

    stato["rotto"] = False
    d._dimentica("k")                    # ritenta dopo la correzione
    d._export_all()
    assert (tmp_path / "k.tex").read_text(encoding="utf-8") == "contenuto valido"


def test_b12_batch_misto_conta_solo_i_riusciti(tmp_path, dialogo_senza_finestre):
    items = [(f"k{i}", "", f"k{i}.svg",
              (lambda i=i: f"<svg>{i}</svg>") if i % 2 == 0
              else (lambda: (_ for _ in ()).throw(RuntimeError("x"))))
             for i in range(4)]
    d = _FintoDialogo(items, tmp_path)
    d._export_all()
    creati = sorted(p.name for p in tmp_path.iterdir())
    assert creati == ["k0.svg", "k2.svg"]
    completato = [m for t, _ti, m in dialogo_senza_finestre if t == "info"][0]
    assert "2" in completato


# ─────────────────────────────── R02 ────────────────────────────────────────

def test_r02_excel_fallito_preserva_il_file_precedente(tmp_path, monkeypatch):
    pytest.importorskip("openpyxl")
    import openpyxl

    dest = tmp_path / "analisi.xlsx"
    scrivi_excel([{"perm_str": "[0]", "simboliche": ["prima"], "n_sim": 1}],
                 str(dest))
    prima = dest.read_bytes()

    def save_rotto(self, f):
        f.write(b"meta' file")
        raise RuntimeError("disco pieno")

    monkeypatch.setattr(openpyxl.Workbook, "save", save_rotto)
    with pytest.raises(RuntimeError):
        scrivi_excel([{"perm_str": "[1]", "simboliche": ["dopo"], "n_sim": 1}],
                     str(dest))
    assert dest.read_bytes() == prima, "la destinazione precedente e' intatta"
    assert [p.name for p in tmp_path.iterdir()] == ["analisi.xlsx"], \
        "nessun temporaneo residuo"


def test_r02_excel_grezzo_fallito_preserva_il_file_precedente(
        tmp_path, monkeypatch, senza_dialoghi):
    pytest.importorskip("openpyxl")
    import openpyxl

    dest = tmp_path / "grezzi.xlsx"
    _esporta_excel_grezzo([_risultato(["prima"])], dest, monkeypatch)
    prima = dest.read_bytes()

    def save_rotto(self, f):
        raise RuntimeError("errore durante il salvataggio")

    monkeypatch.setattr(openpyxl.Workbook, "save", save_rotto)
    _esporta_excel_grezzo([_risultato(["dopo"])], dest, monkeypatch)
    assert senza_dialoghi, "l'errore deve essere segnalato, non ignorato"
    assert dest.read_bytes() == prima
    assert [p.name for p in tmp_path.iterdir()] == ["grezzi.xlsx"]


def test_r02_export_dialogo_non_tocca_i_file_altrui(tmp_path, dialogo_senza_finestre):
    vicino = tmp_path / "vicino.txt"
    vicino.write_text("da non toccare", encoding="utf-8")
    dest = tmp_path / "k.svg"
    dest.write_text("versione precedente", encoding="utf-8")

    items = [("k", "", "k.svg", lambda: (_ for _ in ()).throw(RuntimeError("x")))]
    _FintoDialogo(items, tmp_path)._export_all()
    assert dest.read_text(encoding="utf-8") == "versione precedente"
    assert vicino.read_text(encoding="utf-8") == "da non toccare"
    assert sorted(p.name for p in tmp_path.iterdir()) == ["k.svg", "vicino.txt"]


# ─────────────────────────────── R03 ────────────────────────────────────────

def _pdf_di_prova(path, pagine=2):
    pytest.importorskip("reportlab")
    from reportlab.pdfgen import canvas

    c = canvas.Canvas(str(path))
    for i in range(pagine):
        c.drawString(40, 40, f"pagina {i}")
        c.showPage()
    c.save()
    return path


def test_r03_temporanei_distinti_nella_cartella_di_destinazione(tmp_path):
    dest = tmp_path / "out.pdf"
    dest.write_bytes(b"%PDF-1.4\n")
    a = pdfmerge._temporaneo(str(dest))
    b = pdfmerge._temporaneo(str(dest))
    try:
        assert a != b
        assert pathlib.Path(a).parent == pathlib.Path(b).parent == tmp_path
        pdfmerge._pulisci(a)
        assert pathlib.Path(b).exists(), "la pulizia di uno non tocca l'altro"
    finally:
        pdfmerge._pulisci(a)
        pdfmerge._pulisci(b)


def test_r03_due_deduplicazioni_sulla_stessa_destinazione(tmp_path, monkeypatch):
    """Concorrenza simulata in modo deterministico: la seconda parte dentro la prima."""
    pytest.importorskip("pypdf")
    dest = _pdf_di_prova(tmp_path / "out.pdf")
    temporanei = []
    originale_temp = pdfmerge._temporaneo
    originale_sost = pdfmerge._sostituisci
    annidata = {"fatta": False, "esito": None}

    def spia_temp(path):
        t = originale_temp(path)
        temporanei.append(t)
        return t

    def sostituisci_intrecciato(tmp, path):
        # mentre la prima deduplicazione ha il suo temporaneo pronto ma non
        # ancora pubblicato, ne parte una seconda sulla stessa destinazione
        if not annidata["fatta"]:
            annidata["fatta"] = True
            assert pathlib.Path(tmp).exists()
            annidata["esito"] = pdfmerge._dedup_pypdf(str(path))
            assert pathlib.Path(tmp).exists(), \
                "la seconda deduplicazione ha distrutto il temporaneo della prima"
        return originale_sost(tmp, path)

    monkeypatch.setattr(pdfmerge, "_temporaneo", spia_temp)
    monkeypatch.setattr(pdfmerge, "_sostituisci", sostituisci_intrecciato)

    assert pdfmerge._dedup_pypdf(str(dest)) is True
    assert annidata["esito"] is True
    assert len(set(temporanei)) == len(temporanei) == 2

    from pypdf import PdfReader
    assert len(PdfReader(str(dest)).pages) == 2
    assert sorted(p.name for p in tmp_path.iterdir()) == ["out.pdf"], \
        "nessun temporaneo sopravvissuto"


# ─────────────────────────────── M05 ────────────────────────────────────────

@pytest.mark.parametrize("nome", ["con spazi.html", "àccénti.html",
                                  "cancelletto #1.html", "misto à #2 b.html"])
def test_m05_uri_del_browser_riporta_al_file(nome, tmp_path, monkeypatch,
                                             senza_dialoghi):
    """L'URL aperto dal browser deve ri-tradursi esattamente nel file scritto."""
    import webbrowser

    dest = tmp_path / nome
    aperti = []
    monkeypatch.setattr(webbrowser, "open", lambda url: aperti.append(url))
    monkeypatch.setattr(at, "filedialog", types.SimpleNamespace(
        asksaveasfilename=lambda **k: str(dest)))
    tab = _Analisi(risultati=[_risultato(["(SCD_U x SCD_U x SCD_U)"])])
    tab._export_analisi_html()

    assert dest.exists() and dest.read_text(encoding="utf-8").startswith("<!DOCTYPE")
    assert len(aperti) == 1
    url = aperti[0]
    assert url.startswith("file://") and " " not in url
    ricavato = urllib.request.url2pathname(urllib.parse.urlparse(url).path)
    assert pathlib.Path(ricavato) == dest.resolve()


def test_m05_nessun_uri_costruito_a_mano():
    """Nessun modulo del package deve piu' concatenare 'file://' + percorso."""
    sorgenti = list((RADICE / "gioco27").rglob("*.py"))
    colpevoli = [p.relative_to(RADICE).as_posix() for p in sorgenti
                 if 'f"file://' in p.read_text(encoding="utf-8")]
    assert colpevoli == []
