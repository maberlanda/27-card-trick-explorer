"""
Tab "📊 Analisi" — molteplicità delle permutazioni e relativi export.

Mixin per App: fornisce _build_analisi_tab(), i callback di analisi e
gli export (CSV, Excel, HTML, dati grezzi).
Estratto da app.py (v2.8.0) senza modifiche funzionali.
"""
import os
import tkinter as tk
from tkinter import ttk, messagebox, filedialog

from ..core.algebra import (analizza_righe, analizza_csv, _prep_explorer_expr,
                             scrivi_output, scrivi_excel)
from ..core.combinations import iter_combinations_ex, count_combinations_ex
from ..core.permutations import make_csv_row
from .common import configure_matrix_tags, insert_colored, EtaEstimator, run_in_thread
from .i18n import tr
from .i18n import get_language


class AnalysisTabMixin:
    """Metodi del tab Analisi. Richiede gli attributi/metodi di App:
    _get_filters(), progress, _explorer_entry, _explorer_calc(),
    _switch_to_explorer()."""

    def _build_analisi_tab(self, parent):
        outer = ttk.Frame(parent, padding=10)
        outer.rowconfigure(2, weight=1)
        outer.columnconfigure(0, weight=1)

        hdr = ttk.Frame(outer)
        hdr.grid(row=0, column=0, sticky="ew", pady=(0, 8))
        ttk.Label(hdr,
                  text=f"📊  {tr('analysis.title')}",
                  font=("Segoe UI", 13, "bold"),
                  foreground="#1a5276").pack(side="left")
        ttk.Label(hdr,
                  text=f"   {tr('analysis.subtitle')}",
                  font=("Segoe UI", 10, "italic"),
                  foreground="#555").pack(side="left")

        cmd = ttk.Frame(outer)
        cmd.grid(row=1, column=0, sticky="ew", pady=(0, 6))
        ttk.Button(cmd, text=f"🔄  {tr('analysis.generate')}",
                   style="Action.TButton",
                   command=self._run_analisi).pack(side="left", padx=(0, 6))
        self._analisi_exp_mb = tk.Menubutton(cmd, text=tr("analysis.export"),
                                              relief="raised", state="disabled")
        exp_menu2 = tk.Menu(self._analisi_exp_mb, tearoff=0)
        exp_menu2.add_command(label=f"📊  {tr('analysis.menu.summary_csv')}",
                              command=self._export_analisi_csv)
        exp_menu2.add_command(label=f"📗  {tr('analysis.menu.summary_excel')}",
                              command=self._export_analisi_excel)
        exp_menu2.add_command(label=f"🌐  {tr('analysis.menu.summary_html')}",
                              command=self._export_analisi_html)
        exp_menu2.add_separator()
        exp_menu2.add_command(label=f"📋  {tr('analysis.menu.raw_csv')}",
                              command=self._export_analisi_raw_csv)
        exp_menu2.add_command(label=f"📗  {tr('analysis.menu.raw_excel')}",
                              command=self._export_analisi_raw_excel)
        self._analisi_exp_mb["menu"] = exp_menu2
        self._analisi_exp_mb.pack(side="left", padx=4)
        ttk.Button(cmd, text=f"🔬  {tr('analysis.open_explorer')}",
                   style="Preset.TButton",
                   command=self._analisi_to_explorer).pack(side="left", padx=10)
        self._analisi_status = tk.StringVar(
            value=tr("analysis.status.prompt"))
        ttk.Label(cmd, textvariable=self._analisi_status,
                  font=("Segoe UI", 10, "italic"), foreground="#555",
                  wraplength=500).pack(side="left", padx=8)

        # PanedWindow verticale: lista sopra, dettaglio sotto
        paned = ttk.PanedWindow(outer, orient="vertical")
        paned.grid(row=2, column=0, sticky="nsew")

        # ── Pannello superiore: Treeview ──
        tv_frame = ttk.Frame(paned)
        tv_frame.rowconfigure(0, weight=1)
        tv_frame.columnconfigure(0, weight=1)
        paned.add(tv_frame, weight=3)

        cols = ("n_sim", "perm", "simbolica_0")
        self._analisi_tv = ttk.Treeview(
            tv_frame, columns=cols, show="headings", selectmode="browse")
        self._analisi_tv.heading(
            "n_sim", text=tr("analysis.column.multiplicity"))
        self._analisi_tv.heading(
            "perm", text=tr("analysis.column.permutation"))
        self._analisi_tv.heading(
            "simbolica_0", text=tr("analysis.column.first_symbolic"))
        self._analisi_tv.column("n_sim",       width=70,  anchor="center", stretch=False)
        self._analisi_tv.column("perm",        width=340, anchor="w")
        self._analisi_tv.column("simbolica_0", width=600, anchor="w")
        vsb = ttk.Scrollbar(tv_frame, orient="vertical",   command=self._analisi_tv.yview)
        hsb = ttk.Scrollbar(tv_frame, orient="horizontal", command=self._analisi_tv.xview)
        self._analisi_tv.configure(yscrollcommand=vsb.set, xscrollcommand=hsb.set)
        self._analisi_tv.grid(row=0, column=0, sticky="nsew")
        vsb.grid(row=0, column=1, sticky="ns")
        hsb.grid(row=1, column=0, sticky="ew")
        self._analisi_tv.bind("<Double-1>", lambda e: self._analisi_to_explorer())
        self._analisi_tv.bind("<<TreeviewSelect>>", self._analisi_show_detail)
        self._analisi_tv.tag_configure("odd",  background="#F0F4FA")
        self._analisi_tv.tag_configure("even", background="#FFFFFF")
        self._analisi_tv.tag_configure("top",  background="#EAF4FB",
                                               font=("Segoe UI", 10, "bold"))

        # ── Pannello inferiore: dettaglio simboliche ──
        det_frame = ttk.Frame(paned)
        det_frame.rowconfigure(1, weight=1)
        det_frame.columnconfigure(0, weight=1)
        paned.add(det_frame, weight=2)

        det_hdr = ttk.Frame(det_frame)
        det_hdr.grid(row=0, column=0, sticky="ew", pady=(4, 2))
        ttk.Label(det_hdr,
                  text=f"🔍  {tr('analysis.detail.title')}",
                  font=("Segoe UI", 11, "bold"),
                  foreground="#1a5276").pack(side="left")
        ttk.Label(det_hdr,
                  text=f"   {tr('analysis.detail.subtitle')}",
                  font=("Segoe UI", 9, "italic"),
                  foreground="#777").pack(side="left")

        det_txt_frame = ttk.Frame(det_frame)
        det_txt_frame.grid(row=1, column=0, sticky="nsew")
        det_txt_frame.rowconfigure(0, weight=1)
        det_txt_frame.columnconfigure(0, weight=1)

        self._analisi_detail_text = tk.Text(
            det_txt_frame,
            font=("Courier New", 9),
            state="disabled",
            wrap="none",
            background="#FAFBFC",
            relief="flat",
            borderwidth=1,
        )
        dt_vsb = ttk.Scrollbar(det_txt_frame, orient="vertical",
                                command=self._analisi_detail_text.yview)
        dt_hsb = ttk.Scrollbar(det_txt_frame, orient="horizontal",
                                command=self._analisi_detail_text.xview)
        self._analisi_detail_text.configure(
            yscrollcommand=dt_vsb.set, xscrollcommand=dt_hsb.set)
        self._analisi_detail_text.grid(row=0, column=0, sticky="nsew")
        dt_vsb.grid(row=0, column=1, sticky="ns")
        dt_hsb.grid(row=1, column=0, sticky="ew")

        # tag per colorazione
        self._analisi_detail_text.tag_configure(
            "title",  font=("Courier New", 9, "bold"), foreground="#1a3a5c")
        self._analisi_detail_text.tag_configure(
            "idx",    font=("Courier New", 9, "bold"), foreground="#1a5276")
        self._analisi_detail_text.tag_configure(
            "tsim",   font=("Courier New", 9),         foreground="#333333")
        self._analisi_detail_text.tag_configure(
            "stage",  font=("Courier New", 9),         foreground="#555555")
        self._analisi_detail_text.tag_configure(
            "pval",   font=("Courier New", 9, "bold"), foreground="#1a6e2f")
        self._analisi_detail_text.tag_configure(
            "jval",   font=("Courier New", 9, "bold"), foreground="#7b3d00")
        self._analisi_detail_text.tag_configure(
            "hint",   font=("Courier New", 9, "italic"), foreground="#AAAAAA")

        configure_matrix_tags(self._analisi_detail_text)
        _hint = f"  {tr('analysis.detail.hint')}"
        self._analisi_detail_text.configure(state="normal")
        self._analisi_detail_text.insert("1.0", _hint, "hint")
        self._analisi_detail_text.configure(state="disabled")

        self._analisi_risultati = []
        self._analisi_righe_raw = []
        return outer

    def _reset_analisi(self):
        """Riporta il tab Analisi allo stato iniziale (per «Reset tutto»)."""
        self._analisi_risultati = []
        self._analisi_righe_raw = []
        tv = getattr(self, "_analisi_tv", None)
        if tv is not None:
            tv.delete(*tv.get_children())
        txt = getattr(self, "_analisi_detail_text", None)
        if txt is not None:
            txt.configure(state="normal")
            txt.delete("1.0", "end")
            txt.insert("1.0", f"  {tr('analysis.detail.hint')}", "hint")
            txt.configure(state="disabled")
        if hasattr(self, "_analisi_exp_mb"):
            self._analisi_exp_mb.configure(state="disabled")
        self._analisi_status.set(tr("analysis.status.prompt"))

    def _run_analisi(self):
        filters = self._get_filters()
        n = count_combinations_ex(filters)
        if n == 0:
            messagebox.showwarning(tr("analysis.no_combinations_title"),
                                   tr("analysis.no_combinations"))
            return
        if n > 50_000:
            ok = messagebox.askyesno(
                tr("analysis.large_title"),
                tr("analysis.large_confirmation", count=f"{n:,}"))
            if not ok:
                return
        self._analisi_status.set(
            tr("analysis.status.generating", count=f"{n:,}"))
        self.progress["maximum"] = n
        self.progress["value"]   = 0
        self.update_idletasks()

        def job():
            _eta = EtaEstimator()
            righe = []
            for i, params in enumerate(iter_combinations_ex(filters), 1):
                if getattr(self, "_closing", False):
                    return
                rd = make_csv_row(i, params)
                righe.append({
                    "Stage0": rd[1], "Stage1": rd[2], "Stage2": rd[3],
                    "A0": rd[4], "A1": rd[5], "A2": rd[6],
                    "T_simbolica": rd[7], "T_permutazione": rd[8],
                })
                if i % 200 == 0:
                    self._ui(lambda v=i: (
                        self.progress.__setitem__("value", v),
                        self._analisi_status.set(tr(
                            "analysis.status.generation_progress",
                            done=f"{v:,}", total=f"{n:,}", eta=_eta.text(v, n)))))
            if getattr(self, "_closing", False):
                return
            self._analisi_righe_raw = righe
            self._ui(lambda: self._analisi_status.set(tr(
                "analysis.status.analyzing_rows", count=f"{len(righe):,}")))
            risultati = analizza_righe(righe)
            if getattr(self, "_closing", False):
                return
            self._analisi_risultati = risultati
            self._ui(lambda: self._analisi_populate(risultati, n))

        run_in_thread(self, job, error_title=tr("analysis.error_title"),
                      on_error=lambda e: self._analisi_status.set(
                          tr("analysis.status.error")))

    def _analisi_load_csv(self):
        """Carica un CSV COMBINAZIONI esterno e avvia l'analisi Stage-level."""
        path = filedialog.askopenfilename(
            title=tr("analysis.open_csv"),
            filetypes=[("CSV", "*.csv"),
                       (tr("common.filetype_all"), "*.*")])
        if not path:
            return
        self._analisi_status.set(tr(
            "analysis.status.reading_csv", filename=os.path.basename(path)))
        self.update_idletasks()

        def job():
            risultati = analizza_csv(path)
            if getattr(self, "_closing", False):
                return
            n = sum(r["n_sim"] for r in risultati)
            self._analisi_risultati = risultati
            self._ui(lambda: self._analisi_populate(risultati, n))

        run_in_thread(self, job, error_title=tr("analysis.csv_read_error_title"),
                      on_error=lambda e: self._analisi_status.set(
                          tr("analysis.status.error")))


    def _analisi_csv_pipeline(self):
        """Pipeline completa: carica COMBINAZIONI CSV → analisi → salva CSV + Excel.
        Replica esattamente il comportamento di analisi_sequenze.py."""
        inp = filedialog.askopenfilename(
            title=tr("analysis.open_csv"),
            filetypes=[("CSV", "*.csv"),
                       (tr("common.filetype_all"), "*.*")])
        if not inp:
            return

        import os as _os
        base, _ = _os.path.splitext(inp)
        out_csv  = filedialog.asksaveasfilename(
            title=tr("analysis.save_pipeline"),
            initialfile=_os.path.basename(base) + "_analisi.csv",
            defaultextension=".csv",
            filetypes=[("CSV", "*.csv")])
        if not out_csv:
            return
        out_xlsx = _os.path.splitext(out_csv)[0] + ".xlsx"

        self._analisi_status.set(tr("analysis.status.running"))
        self.update_idletasks()

        def job():
            risultati = analizza_csv(inp)
            if getattr(self, "_closing", False):
                return
            scrivi_output(risultati, out_csv)
            if getattr(self, "_closing", False):
                return
            scrivi_excel(risultati, out_xlsx)
            if getattr(self, "_closing", False):
                return
            n_perm = len(risultati)
            n_tot  = sum(r["n_sim"] for r in risultati)
            self._analisi_risultati = risultati
            self._ui(lambda: self._analisi_populate(risultati, n_tot))
            msg = tr(
                "analysis.completed_summary", permutations=f"{n_perm:,}",
                sequences=f"{n_tot:,}", csv=_os.path.basename(out_csv),
                excel=_os.path.basename(out_xlsx))
            self._ui(lambda m=msg: messagebox.showinfo(
                tr("analysis.completed_title"), m))

        run_in_thread(self, job, error_title=tr("error.generic"),
                      on_error=lambda e: self._analisi_status.set(
                          tr("analysis.status.error")))


    def _analisi_populate(self, risultati, n_tot):
        self.progress["value"] = n_tot
        tv = self._analisi_tv
        tv.delete(*tv.get_children())
        max_m = risultati[0]["n_sim"] if risultati else 0
        for idx, r in enumerate(risultati):
            tag = "top" if r["n_sim"] == max_m else ("odd" if idx % 2 else "even")
            sim0 = r["simboliche"][0] if r["simboliche"] else ""
            tv.insert("", "end", values=(r["n_sim"], r["perm_str"], sim0),
                      tags=(tag,), iid=str(idx))
        n_dist = len(risultati)
        min_m  = risultati[-1]["n_sim"] if risultati else 0
        if hasattr(self, "_analisi_exp_mb"):
            self._analisi_exp_mb.configure(state="normal")
        self._analisi_status.set("✓  " + tr(
            "analysis.status.summary", combinations=f"{n_tot:,}",
            permutations=f"{n_dist:,}", minimum=min_m, maximum=max_m))

    def _analisi_show_detail(self, event=None):
        """Mostra nel pannello inferiore tutte le sequenze Stage della riga selezionata,
        con i valori P e J di ciascuno stadio estratti dalla Stage string."""
        txt = self._analisi_detail_text
        txt.configure(state="normal")
        txt.delete("1.0", "end")

        sel = self._analisi_tv.selection()
        if not sel or not self._analisi_risultati:
            txt.insert("end", f"  {tr('analysis.detail.select_row')}", "hint")
            txt.configure(state="disabled")
            return

        idx = int(sel[0])
        r   = self._analisi_risultati[idx]
        simboliche = r["simboliche"]   # lista ordinata di Stage string
        n   = len(simboliche)

        def parse_stage_entry(stage_str):
            """Estrae Stage0/1/2 dalla Stage string 'T = [s3] o [s2] o [s1]'.
            Restituisce lista [s3, s2, s1] o None se non parsabile."""
            prefix = "T = ["
            if not stage_str.startswith(prefix):
                return None
            inner = stage_str[len(prefix):-1]   # rimuove 'T = [' e ']' finale
            parts = inner.split("] o [")
            return parts if len(parts) == 3 else None

        def parse_stage(stage_str):
            """Estrae (P_label, J_label) da 'P_name o MSC o J_name' (Stage singolo)."""
            parts = stage_str.split(" o MSC o ", 1)
            if len(parts) == 2:
                return parts[0].strip(), parts[1].strip()
            parts = stage_str.split(" ∘ msc ∘ ", 1)
            if len(parts) == 2:
                return parts[0].strip(), parts[1].strip()
            return stage_str.strip(), ""

        multiplicity_key = ("analysis.detail.multiplicity.one" if n == 1
                            else "analysis.detail.multiplicity.many")
        txt.insert("end", f"  {tr(multiplicity_key, count=n)}\n", "title")
        txt.insert("end", f"  {tr('analysis.detail.permutation', permutation=r['perm_str'])}\n\n",
                   "tsim")

        # intestazione colonne
        hdr = (f"  {'#':>2}  "
               f"{'Stage0 — P':25s}  {'J':25s}    "
               f"{'Stage1 — P':25s}  {'J':25s}    "
               f"{'Stage2 — P':25s}  {'J':25s}\n")
        sep = "  " + "─" * max(len(hdr) - 4, 60) + "\n"
        txt.insert("end", hdr,  "title")
        txt.insert("end", sep,  "title")

        # rimuovi tag-link precedenti
        for tag in txt.tag_names():
            if tag.startswith("link_"):
                txt.tag_delete(tag)

        for i, ts in enumerate(simboliche, 1):
            parts = parse_stage_entry(ts)   # [s3, s2, s1]
            if parts:
                # ordine visivo: Stage0, Stage1, Stage2
                cells = [parse_stage(parts[2]), parse_stage(parts[1]), parse_stage(parts[0])]
            else:
                cells = [("", ""), ("", ""), ("", "")]
            row_txt = f"  {i:>2}  "
            txt.insert("end", row_txt, "idx")
            for si, (p_lbl, j_lbl) in enumerate(cells):
                p_str = f"{p_lbl:<25s}"
                j_str = f"{j_lbl:<25s}"
                insert_colored(txt, p_str, "pval")
                txt.insert("end", "  ", "stage")
                insert_colored(txt, j_str, "jval")
                if si < 2:
                    txt.insert("end", "    ", "stage")
            txt.insert("end", "\n", "stage")
            # riga secondaria: Stage string cliccabile → Explorer.
            # Dalla v2.8.4 viene inviata la notazione di gioco COMPLETA
            # (P o MSC o J per ogni turno, con R_U/I_3 visibili), che
            # l'Explorer ora accetta nativamente: ciò che vedi qui è
            # esattamente ciò che viene calcolato.
            link_tag = f"link_{i}"
            txt.tag_configure(link_tag, foreground="#0a5fa8", underline=True)
            txt.tag_bind(link_tag, "<Enter>",
                         lambda e, w=txt: w.configure(cursor="hand2"))
            txt.tag_bind(link_tag, "<Leave>",
                         lambda e, w=txt: w.configure(cursor=""))
            txt.tag_bind(link_tag, "<Button-1>",
                         lambda e, t=ts: self._analisi_send_to_explorer(t))
            lbl = f"      → {ts}  ({tr('analysis.detail.open_explorer')})"
            insert_colored(txt, lbl, (link_tag, "tsim"))
            txt.insert("end", "\n\n", "tsim")

        # nota in fondo
        txt.insert("end", f"  {tr('analysis.detail.footer_hint')}\n", "hint")
        txt.configure(state="disabled")

    def _analisi_send_to_explorer(self, t_sim):
        """Invia una specifica T_simbolica (dal dettaglio) all'Explorer."""
        expr = _prep_explorer_expr(t_sim)
        self._explorer_entry.delete("1.0", "end")
        self._explorer_entry.insert("1.0", expr)
        self._switch_to_explorer()
        self._explorer_calc()

    def _analisi_to_explorer(self):
        sel = self._analisi_tv.selection()
        if not sel:
            messagebox.showinfo(tr("analysis.selection_title"),
                                tr("analysis.selection_required"))
            return
        idx   = int(sel[0])
        r     = self._analisi_risultati[idx]
        stage0 = r["simboliche"][0] if r["simboliche"] else ""
        if not stage0:
            messagebox.showinfo("Explorer", tr("analysis.no_symbolic"))
            return
        # Dalla v2.8.4 si invia direttamente la Stage string (P e J separati)
        expr = _prep_explorer_expr(stage0)
        self._explorer_entry.delete("1.0", "end")
        self._explorer_entry.insert("1.0", expr)
        self._switch_to_explorer()
        self._explorer_calc()



    def _export_analisi_csv(self):
        if not self._analisi_risultati:
            messagebox.showinfo(tr("analysis.no_data_title"),
                                tr("analysis.no_data"))
            return
        path = filedialog.asksaveasfilename(
            defaultextension=".csv", filetypes=[("CSV", "*.csv")],
            title=tr("analysis.save_summary_csv"))
        if not path:
            return
        try:
            scrivi_output(self._analisi_risultati, path)
            self._analisi_status.set("✓  " + tr(
                "analysis.status.exported", format="CSV",
                filename=os.path.basename(path)))
        except Exception as e:
            messagebox.showerror(tr("error.generic"), str(e))

    def _export_analisi_excel(self):
        if not self._analisi_risultati:
            messagebox.showinfo(tr("analysis.no_data_title"),
                                tr("analysis.no_data"))
            return
        path = filedialog.asksaveasfilename(
            defaultextension=".xlsx", filetypes=[("Excel", "*.xlsx")],
            title=tr("analysis.save_summary_excel"))
        if not path:
            return
        try:
            scrivi_excel(self._analisi_risultati, path)
            self._analisi_status.set("✓  " + tr(
                "analysis.status.exported", format="Excel",
                filename=os.path.basename(path)))
        except Exception as e:
            messagebox.showerror(tr("error.generic"), str(e))


    def _export_analisi_html(self):
        """Esporta l'analisi molteplicità in HTML leggibile (una simbolica per riga)."""
        if not self._analisi_risultati:
            messagebox.showinfo(tr("analysis.no_data_title"),
                                tr("analysis.no_data"))
            return
        import html as _h
        path = filedialog.asksaveasfilename(
            defaultextension=".html", filetypes=[("HTML", "*.html")],
            title=tr("analysis.save_summary_html"))
        if not path:
            return
        try:
            rows = ""
            for r in self._analisi_risultati:
                syms_html = "".join(
                    f"<div class='sym'>{_h.escape(s)}</div>"
                    for s in r.get("simboliche", [])
                )
                rows += (
                    f"<tr>"
                    f"<td class='molt'>{r['n_sim']}</td>"
                    f"<td class='perm'><code>{_h.escape(r['perm_str'])}</code></td>"
                    f"<td class='syms'>{syms_html}</td>"
                    f"</tr>\n"
                )
            css = (
                "body{font-family:'Segoe UI',sans-serif;padding:24px;background:#f9f9f9}"
                "h2{color:#1a5276;margin-bottom:12px}"
                "table{border-collapse:collapse;width:100%;background:#fff;"
                "      box-shadow:0 1px 4px rgba(0,0,0,.1)}"
                "th{background:#1a5276;color:#fff;padding:8px 12px;"
                "   font-family:'Courier New',monospace;font-size:.85em;text-align:left}"
                "td{border:1px solid #ddd;padding:6px 10px;vertical-align:top;"
                "   font-family:'Courier New',monospace;font-size:.82em}"
                "tr:nth-child(even) td{background:#f4f8ff}"
                "td.molt{text-align:right;font-weight:bold;color:#1a5276;width:60px}"
                "td.perm{white-space:nowrap;width:400px}"
                "div.sym{padding:1px 0;border-bottom:1px dotted #ddd;white-space:pre-wrap;"
                "        word-break:break-all}"
                "div.sym:last-child{border-bottom:none}"
            )
            n_perm = len(self._analisi_risultati)
            n_tot  = sum(r["n_sim"] for r in self._analisi_risultati)
            html_str = (
                f"<!DOCTYPE html><html lang='{get_language()}'><head><meta charset='UTF-8'>"
                f"<title>{tr('export.document.analysis.title')}</title>"
                f"<style>{css}</style></head><body>"
                f"<h2>{tr('export.document.analysis.heading')}</h2>"
                f"<p>{tr('export.document.analysis.summary', combinations=n_tot, permutations=n_perm)}</p>"
                f"<table><thead><tr>"
                f"<th>{tr('export.document.analysis.multiplicity')}</th><th>T_permutazione [0..26]</th>"
                f"<th>{tr('export.document.analysis.distinct_stages')}</th>"
                f"</tr></thead><tbody>"
                f"{rows}</tbody></table></body></html>"
            )
            with open(path, "w", encoding="utf-8") as f:
                f.write(html_str)
            self._analisi_status.set("✓  " + tr(
                "analysis.status.exported", format="HTML",
                filename=os.path.basename(path)))
            import webbrowser
            webbrowser.open(f"file://{path}")
        except Exception as e:
            messagebox.showerror(tr("error.generic"), str(e))

    def _export_analisi_raw_csv(self):
        """Esporta i dati grezzi (una riga per combinazione) in CSV con sep=;"""
        if not self._analisi_righe_raw:
            messagebox.showinfo(tr("analysis.no_data_title"),
                                tr("analysis.no_data"))
            return
        path = filedialog.asksaveasfilename(
            defaultextension=".csv", filetypes=[("CSV", "*.csv")],
            title=tr("analysis.save_raw_csv"))
        if not path:
            return
        try:
            import csv as _csv
            cols = ["#", "Stage0", "Stage1", "Stage2",
                    "A0", "A1", "A2", "T_simbolica", "T_permutazione"]
            with open(path, "w", newline="", encoding="utf-8-sig") as f:
                w = _csv.writer(f, delimiter=";", quotechar='"',
                                quoting=_csv.QUOTE_ALL)
                w.writerow(cols)
                for i, r in enumerate(self._analisi_righe_raw, 1):
                    w.writerow([
                        i,
                        r.get("Stage0", ""), r.get("Stage1", ""), r.get("Stage2", ""),
                        r.get("A0", ""),     r.get("A1", ""),     r.get("A2", ""),
                        r.get("T_simbolica", ""), r.get("T_permutazione", ""),
                    ])
            self._analisi_status.set("✓  " + tr(
                "analysis.status.raw_exported", format="CSV",
                filename=os.path.basename(path)))
        except Exception as e:
            messagebox.showerror(tr("error.generic"), str(e))

    def _export_analisi_raw_excel(self):
        """Esporta i dati grezzi (una riga per combinazione) in Excel."""
        if not self._analisi_righe_raw:
            messagebox.showinfo(tr("analysis.no_data_title"),
                                tr("analysis.no_data"))
            return
        path = filedialog.asksaveasfilename(
            defaultextension=".xlsx", filetypes=[("Excel", "*.xlsx")],
            title=tr("analysis.save_raw_excel"))
        if not path:
            return
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment
            wb = openpyxl.Workbook()

            # ── Foglio 1: dati grezzi ──────────────────────────────────────
            ws1 = wb.active
            ws1.title = "Dati grezzi"
            cols = ["#", "Stage0", "Stage1", "Stage2",
                    "A0", "A1", "A2", "T_simbolica", "T_permutazione"]
            hdr_fill  = PatternFill("solid", fgColor="1A5276")
            hdr_font  = Font(bold=True, color="FFFFFF", name="Courier New", size=9)
            data_font = Font(name="Courier New", size=9)
            wrap_al   = Alignment(wrap_text=True, vertical="top")
            for ci, col in enumerate(cols, 1):
                cell = ws1.cell(row=1, column=ci, value=col)
                cell.fill = hdr_fill
                cell.font = hdr_font
                cell.alignment = Alignment(horizontal="center", vertical="center")
            for ri, r in enumerate(self._analisi_righe_raw, 2):
                vals = [
                    ri - 1,
                    r.get("Stage0", ""), r.get("Stage1", ""), r.get("Stage2", ""),
                    r.get("A0", ""),     r.get("A1", ""),     r.get("A2", ""),
                    r.get("T_simbolica", ""), r.get("T_permutazione", ""),
                ]
                for ci, v in enumerate(vals, 1):
                    cell = ws1.cell(row=ri, column=ci, value=v)
                    cell.font = data_font
                    cell.alignment = wrap_al
            # larghezze colonne
            for ci, w in enumerate([6, 40, 40, 40, 28, 28, 32, 60, 60], 1):
                ws1.column_dimensions[
                    openpyxl.utils.get_column_letter(ci)].width = w
            ws1.freeze_panes = "A2"

            # ── Foglio 2: analisi molteplicità ────────────────────────────
            ws2 = wb.create_sheet("Analisi molteplicità")
            h2 = ["Molteplicità", "T_permutazione", "Sequenze Stage distinte"]
            for ci, col in enumerate(h2, 1):
                cell = ws2.cell(row=1, column=ci, value=col)
                cell.fill = hdr_fill
                cell.font = hdr_font
                cell.alignment = Alignment(horizontal="center", vertical="center")
            for ri, r in enumerate(self._analisi_risultati, 2):
                syms = "\n".join(r.get("simboliche", []))
                ws2.cell(row=ri, column=1, value=r["n_sim"]).font = Font(bold=True, size=10)
                cell_p = ws2.cell(row=ri, column=2, value=r["perm_str"])
                cell_p.font = Font(name="Courier New", size=9)
                cell_s = ws2.cell(row=ri, column=3, value=syms)
                cell_s.font = Font(name="Courier New", size=9)
                cell_s.alignment = Alignment(wrap_text=True, vertical="top")
                ws2.row_dimensions[ri].height = max(15, 14 * r["n_sim"])
            for ci, w in enumerate([12, 55, 100], 1):
                ws2.column_dimensions[
                    openpyxl.utils.get_column_letter(ci)].width = w
            ws2.freeze_panes = "A2"

            wb.save(path)
            self._analisi_status.set("✓  " + tr(
                "analysis.status.raw_exported", format="Excel",
                filename=os.path.basename(path)))
        except Exception as e:
            messagebox.showerror(tr("error.generic"), str(e))
