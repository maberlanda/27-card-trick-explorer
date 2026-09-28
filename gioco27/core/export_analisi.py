"""Export dei risultati dell'analisi: CSV e Excel.

Compartimento G1. Queste due funzioni vivevano in `core/algebra.py`, il modulo
del linguaggio: scrivevano su disco, conoscevano openpyxl e il formato delle
intestazioni, e tenevano in ostaggio il parser in un file che importava csv e
Workbook. Qui non cambia una riga del loro comportamento — stessi fogli, stesse
intestazioni, stessa pubblicazione atomica di R02, stesso limite di cella di
B01 — cambia solo chi le possiede.

`core.algebra` continua a esporle per i chiamanti storici (si veda il suo
`__getattr__`), ma l'implementazione e' qui e solo qui.
"""

import csv

from ..i18n import tr

__all__ = ["EXCEL_MAX_CELL_CHARS", "scrivi_output", "scrivi_excel"]


def scrivi_output(risultati, output_path):
    from .parallel import atomic_write
    header = [
        tr("export.analysis.csv.header.t_permutation"),
        tr("export.analysis.csv.header.symbolic"),
        tr("export.analysis.csv.header.multiplicity"),
    ]
    with atomic_write(output_path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f, delimiter=";", quotechar='"',
                       quoting=csv.QUOTE_ALL, lineterminator="\n")
        w.writerow(header)
        for r in risultati:
            w.writerow([r["perm_str"], " , ".join(r["simboliche"]), str(r["n_sim"])])


#: Limite di caratteri di UNA cella Excel. Oltre questa soglia openpyxl scrive
#: il valore, ma Excel lo tronca alla riapertura: il dato sparisce senza che
#: nessuno se ne accorga. Chi concatena piu' valori in una cella deve tenerne
#: conto ed esporre altrove il dettaglio completo (B01).
EXCEL_MAX_CELL_CHARS = 32767


def scrivi_excel(risultati, output_path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    wb       = Workbook()
    HDR_FILL = PatternFill("solid", fgColor="1F4E79")
    HDR_FONT = Font(bold=True, color="FFFFFF", size=10)
    ALT_FILL = PatternFill("solid", fgColor="D6E4F0")
    BORDER   = Border(bottom=Side(style="thin", color="AAAAAA"),
                      right=Side(style="thin",  color="AAAAAA"))
    def _fmt(ws, i, nc):
        fill = ALT_FILL if i % 2 == 0 else None
        for c in range(1, nc+1):
            cell = ws.cell(i+1, c)
            if fill: cell.fill = fill
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = BORDER
    def _hdr(ws, hdrs):
        ws.append(hdrs)
        for c in range(1, len(hdrs)+1):
            cell = ws.cell(1, c)
            cell.font = HDR_FONT; cell.fill = HDR_FILL
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
    ws1 = wb.active
    ws1.title = "Perm -> Simboliche"
    _hdr(ws1, ["T_permutazione", "T_simboliche_distinte", "n_sim_distinte"])
    for i, r in enumerate(risultati, 1):
        summary = " , ".join(r["simboliche"])
        if len(summary) > EXCEL_MAX_CELL_CHARS:
            summary = tr("export.excel.summary_overflow",
                         count=len(r["simboliche"]))
        ws1.append([r["perm_str"], summary, r["n_sim"]])
        _fmt(ws1, i, 3)
    ws1.column_dimensions["A"].width = 40
    ws1.column_dimensions["B"].width = 120
    ws1.column_dimensions["C"].width = 14
    ws1.freeze_panes = "A2"
    ws1.auto_filter.ref = f"A1:C{len(risultati)+1}"
    ws2 = wb.create_sheet("Simbolica -> Perm")
    _hdr(ws2, ["T_simbolica", "T_permutazione", "n_sim_distinte"])
    righe_inv = sorted(
        [(sim, r["perm_str"], r["n_sim"]) for r in risultati for sim in r["simboliche"]],
        key=lambda x: x[0])
    for i, (sim, perm, n) in enumerate(righe_inv, 1):
        ws2.append([sim, perm, n])
        _fmt(ws2, i, 3)
    ws2.column_dimensions["A"].width = 120
    ws2.column_dimensions["B"].width = 40
    ws2.column_dimensions["C"].width = 14
    ws2.freeze_panes = "A2"
    ws2.auto_filter.ref = f"A1:C{len(righe_inv)+1}"
    # Pubblicazione atomica (R02): finche' il nuovo file non e' completo, al suo
    # posto resta quello precedente, non un .xlsx troncato che Excel rifiuta.
    from .parallel import atomic_write
    with atomic_write(output_path, "wb") as f:
        wb.save(f)
